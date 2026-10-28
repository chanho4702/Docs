from __future__ import annotations

import hashlib
import json
import sys
import zipfile
from pathlib import Path


RUNTIME = Path(r"C:\Users\DMOVE\claude-experiences\.docx_runtime")
sys.path.insert(0, str(RUNTIME))

from docx import Document  # noqa: E402


SOURCE = Path(r"C:\Users\DMOVE\claude-experiences\[시너지파트너즈] - 김찬호 후보자 이력서_에스코어.docx")
OUTPUT = Path(
    r"C:\Users\DMOVE\claude-experiences\[시너지파트너즈] - 김찬호 후보자 이력서_에스코어_무브스아이_상세보강본.docx"
)


def media_hashes(path: Path) -> dict[str, str]:
    with zipfile.ZipFile(path) as archive:
        return {
            name: hashlib.sha256(archive.read(name)).hexdigest()
            for name in archive.namelist()
            if name.startswith("word/media/")
        }


def main() -> None:
    with zipfile.ZipFile(OUTPUT) as archive:
        zip_error = archive.testzip()

    document = Document(OUTPUT)
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    result = {
        "zip_error": zip_error,
        "paragraphs": len(document.paragraphs),
        "tables": len(document.tables),
        "inline_shapes": len(document.inline_shapes),
        "list_paragraphs": sum(
            1
            for paragraph in document.paragraphs
            if paragraph._p.pPr is not None and paragraph._p.pPr.numPr is not None
        ),
        "manual_page_breaks": sum(
            paragraph._p.xml.count('<w:br w:type="page"')
            for paragraph in document.paragraphs
        ),
        "replacement_chars": text.count("\ufffd"),
        "todo_tokens": sum(
            text.upper().count(token) for token in ("TODO", "TBD", "PLACEHOLDER")
        ),
        "media_preserved": media_hashes(SOURCE) == media_hashes(OUTPUT),
        "size_bytes": OUTPUT.stat().st_size,
    }
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
