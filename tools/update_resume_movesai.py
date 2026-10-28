from __future__ import annotations

import copy
import sys
from pathlib import Path


RUNTIME = Path(r"C:\Users\DMOVE\claude-experiences\.docx_runtime")
sys.path.insert(0, str(RUNTIME))

from docx import Document  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402


INPUT = Path(r"C:\Users\DMOVE\claude-experiences\[시너지파트너즈] - 김찬호 후보자 이력서_에스코어.docx")
OUTPUT = Path(r"C:\Users\DMOVE\claude-experiences\[시너지파트너즈] - 김찬호 후보자 이력서_에스코어_무브스아이_상세보강본.docx")


def replace_paragraph_text(paragraph, text: str) -> None:
    """Replace visible text while retaining paragraph and first-run formatting."""
    first_rpr = None
    for run in paragraph.runs:
        if run._r.rPr is not None:
            first_rpr = copy.deepcopy(run._r.rPr)
            break

    for child in list(paragraph._p):
        if child.tag in {qn("w:r"), qn("w:hyperlink"), qn("w:smartTag"), qn("w:sdt")}:
            paragraph._p.remove(child)

    run = paragraph.add_run(text)
    if first_rpr is not None:
        if run._r.rPr is not None:
            run._r.remove(run._r.rPr)
        run._r.insert(0, first_rpr)


def make_paragraph_from_template(template, text: str):
    p = OxmlElement("w:p")
    if template._p.pPr is not None:
        p.append(copy.deepcopy(template._p.pPr))

    r = OxmlElement("w:r")
    rpr = None
    for run in template.runs:
        if run._r.rPr is not None:
            rpr = copy.deepcopy(run._r.rPr)
            break
    if rpr is not None:
        r.append(rpr)

    t = OxmlElement("w:t")
    if text[:1].isspace() or text[-1:].isspace():
        t.set(qn("xml:space"), "preserve")
    t.text = text
    r.append(t)
    p.append(r)
    return p


def is_page_break_paragraph(paragraph) -> bool:
    for br in paragraph._p.iter(qn("w:br")):
        if br.get(qn("w:type")) == "page":
            return True
    return False


def update_summary_tables(doc) -> None:
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                text = cell.text.strip()
                if text.startswith("[회사·업무]") and "Jira Cloud" in text:
                    replace_paragraph_text(
                        cell.paragraphs[0],
                        "[회사·업무] MovesEye(Zabbix 기반 인프라 모니터링)·Moves Workforce(Jira Cloud RMS) 제품 설계, 풀스택 개발, 보안·배포·운영 및 Elastic Stack 구축",
                    )

    replacements = {
        "Backend": "Java 17, Spring Boot 3.3, Spring Security, JPA, Flyway, REST API, Node.js/NestJS",
        "Search & Ops": "Elasticsearch, Zabbix 7, Zabbix JSON-RPC, Kibana, Logstash, Filebeat, Metricbeat, Elastic Agent, Grafana",
        "Infra & Delivery": "Linux, Docker, PostgreSQL, Redis, HashiCorp Vault, Apache, Jenkins, GitLab CI, GitHub Actions",
        "SaaS & Front": "Atlassian Forge, Jira Cloud API, React, TypeScript, Atlaskit, TanStack Query, ECharts, React Flow, MapLibre",
    }
    for table in doc.tables:
        for row in table.rows:
            if len(row.cells) < 2:
                continue
            key = row.cells[0].text.strip()
            if key in replacements:
                replace_paragraph_text(row.cells[1].paragraphs[0], replacements[key])


def insert_core_competency(doc) -> None:
    anchor = next(
        p for p in doc.paragraphs if p.text.startswith("Elasticsearch 클러스터의 master/data/ingest")
    )
    text = (
        "Java 17·Spring Boot 3.3 기반 MovesEye 백엔드 전환과 Zabbix JSON-RPC 수집 엔진, "
        "PostgreSQL·Redis·Flyway·Vault 운영 기반 구축"
    )
    anchor._p.addnext(make_paragraph_from_template(anchor, text))


def replace_movesai_section(doc) -> None:
    paragraphs = doc.paragraphs
    start_idx = next(i for i, p in enumerate(paragraphs) if p.text.strip() == "[담당업무]")
    next_company_idx = next(
        i for i in range(start_idx + 1, len(paragraphs)) if paragraphs[i].text.strip() == "[담당업무]"
    )
    page_break_idx = next(
        i
        for i in range(start_idx + 1, next_company_idx)
        if not paragraphs[i].text.strip() and is_page_break_paragraph(paragraphs[i])
    )

    subhead_template = paragraphs[start_idx]
    bullet_template = next(
        p for p in paragraphs[start_idx + 1 : page_break_idx] if p.text.strip() and p._p.pPr is not None and p._p.pPr.numPr is not None
    )
    prose_template = paragraphs[page_break_idx - 1]
    page_break_anchor = paragraphs[page_break_idx]._p

    content: list[tuple[str, str]] = [
        ("subhead", "[담당업무]"),
        (
            "bullet",
            "MovesEye 인프라 모니터링 제품과 Moves Workforce Jira Cloud RMS의 요구사항 분석, 아키텍처·데이터 모델 설계, 백엔드·프론트엔드 구현, 테스트·배포·운영을 전 범위로 수행",
        ),
        (
            "bullet",
            "Java 17·Spring Boot 3.3·JPA/Flyway·PostgreSQL·Redis 기반 백엔드와 React·TypeScript·TanStack Query 기반 UI 개발, REST/JSON-RPC 계약 및 외부 시스템 연동",
        ),
        (
            "bullet",
            "Atlassian Forge(Node.js 24·arm64·256MB) 서버리스 SaaS에서 Forge SQL/KVS, Jira Cloud API, Async Events·Queue, scheduled/event trigger와 RBAC를 설계·구현",
        ),
        (
            "bullet",
            "Zabbix·Elasticsearch 기반 모니터링/로그 수집·검색·시각화, Linux·Docker·Apache·Vault·CI 환경의 설치·기동·장애 진단과 운영 표준화",
        ),
        (
            "bullet",
            "설계서·API 계약·운영 매뉴얼·QA 증거를 코드와 함께 관리하고 구현자/리뷰어 분리형 AI 하네스로 요구사항-코드-테스트 교차검증 수행",
        ),

        ("subhead", "[주요 프로젝트/실적]  MovesEye(무브스아이) 인프라 통합 모니터링 제품"),
        (
            "bullet",
            "Zabbix가 수집한 CPU·메모리·디스크·네트워크·응답시간·Problem 데이터를 운영 친화적 화면, 권한, 자산·비용·맵 업무로 연결하는 제품을 Spring Boot 3.3.5·Java 17 / React 19 기반으로 개발",
        ),
        (
            "bullet",
            "구 NestJS 백엔드를 Spring Boot 단일 백엔드로 전환하고 auth·user·host·event·setting·Zabbix sync 등 전 기능 모듈과 기존 Controller/DTO·응답 envelope·정렬·시간 직렬화 계약을 보존해 FE 무수정 호환",
        ),
        (
            "bullet",
            "PostgreSQL native enum·jsonb·quoted column·Many-to-Many 정션·native UPSERT를 JPA/Flyway로 이관하고 JWT/bcrypt 상호 호환, Redis 세션, Spring interceptor 적용 범위 차이를 분석·해결",
        ),
        (
            "bullet",
            "Zabbix 7 JSON-RPC client와 host/history/trend/problem/event/user API adapter를 구현하고 chunked body 파싱 실패, inventory 배열/객체 혼용, tags 전체교체, offset 부재 등 실제 API 함정을 진단",
        ),
        (
            "bullet",
            "Gradle app/engine-api/engine-zabbix 멀티모듈과 Metric·Inventory·Alert port를 도입해 Zabbix 구현의 compile 의존을 차단하고, 570개 테스트 및 19개 런타임 응답 그룹의 전환 전후 동등성 검증",
        ),
        (
            "bullet",
            "호스트 그룹 카드/테이블 전환, keyword 검색·서버 정렬, 리소스·알림·가용성 시계열과 상세 drill-down을 구현하고 Zabbix N+1 호출 제거로 목록 응답을 6.4초에서 269ms로 단축",
        ),
        (
            "bullet",
            "호스트별 Zabbix 수집 항목 선택형 커스텀 차트(line/area/bar/scatter/pie/gauge), react-grid-layout 드래그·리사이즈, PostgreSQL jsonb 설정·레이아웃 영속화, history→trend fallback과 단위 정합성 구현",
        ),
        (
            "bullet",
            "호스트 livecheck를 BE OS ping과 Zabbix agent interface availability의 독립 지표로 설계하고, Windows/Linux ProcessBuilder·IP 검증·timeout·CP949 출력과 vfs.fs.size/dependent.size 디스크 key 계열을 처리",
        ),
        (
            "bullet",
            "Assets 코어를 Zabbix/custom 단일 snapshot 모델로 설계하고 사용자 편집 필드와 동기화 소유권 분리, excluded soft delete·복원, 관리자 write 제약을 구현; 실 Zabbix 호스트 75건 동기화 및 핵심 불변식 3종 검증",
        ),
        (
            "bullet",
            "React Flow 기반 draw.io형 Network Map을 신규 개발해 45종 stencil, host/status/problem overlay, diagram_json 문서형 저장, baseUpdatedAt 낙관적 잠금과 view/edit mode를 구현하고 291개 FE 테스트 검증",
        ),
        (
            "bullet",
            "MapLibre 기반 Geo Map에서 Zabbix inventory 좌표를 단일 원천으로 사용하고 핀 drag·직접 입력·미좌표 host 배치, host.update→자산 sync를 연결; N+1·viewport 복원·NaN 좌표·잠김 상태 등 결함 12건 보완",
        ),
        (
            "bullet",
            "호스트 등록 7개 API와 등록 직후 자산 sync/livecheck, 아이템·템플릿 관리 14개 API, 트리거 CRUD 6개 API, host 분류·태그·기기분류·감사로그를 구현하고 관리 IA를 목록→대상→서브탭 drill-down으로 재구성",
        ),
        (
            "bullet",
            "OTP userId 바인딩, 로그인/OTP rate limit, 256bit 미만 JWT secret 부팅 차단, 운영 Swagger off, CORS whitelist와 사용자 열거 방지를 적용하고 access 30분·refresh 8시간 회전·HttpOnly cookie·logout 폐기를 구현",
        ),
        (
            "bullet",
            "개발 secret을 .env에서 HashiCorp Vault KV v2로 이전하고 DB 설정형 scheduler·실행 이력·job별 보관 상한, 자산 비용 대시보드와 영속 Job 기반 CSV 업로드·polling·경합 안전성(CAS/heartbeat/lease) 검수",
        ),
        (
            "bullet",
            "STANDARD/PREMIUM 단일 배포본 라이선스 아키텍처와 GRACE→EXPIRED→ARCHIVED 단계적 차단을 설계하고 RS256 검증·발급 도구·issuer/product 계약 테스트 등 초기 license kernel을 구현",
        ),
        (
            "bullet",
            "사내 VM 온프레미스 배포에서 JWT 부팅 실패, Docker 80→3000 mapping, PostgreSQL 18 dump→16 Flyway 호환, SPA fallback·CORS·HttpOnly cookie·Apache reverse proxy·SELinux 503을 진단하고 복구 절차를 정리",
        ),
        (
            "bullet",
            "Zabbix 관리자 매뉴얼·JSON-RPC API reference·자산/지도/수집 엔진/라이선스 설계서·개발 온보딩과 제품 문서 정본 체계를 구축해 code·OpenAPI·Flyway·route를 근거로 유지",
        ),

        ("subhead", "[주요 프로젝트/실적]  Moves Workforce Jira Cloud 인력·프로젝트 관리 SaaS"),
        (
            "bullet",
            "Atlassian Forge와 React 기반 Jira Cloud 앱을 풀스택 개발하고 32개 resolver registrar, Forge SQL 13개 table, KVS Custom Entity 12종을 중심으로 사용자·팀·프로젝트·워크로그·권한·설정 도메인을 운영",
        ),
        (
            "bullet",
            "Jira 이슈별 워크로그 N+1을 프로젝트 단위 1회 fetch + KVS 10분 cache + p-limit 병렬 전용 fetcher로 전환해 페이지 이동 시 Jira 재호출을 제거하고 캘린더·자원 분석 초기 로딩을 개선",
        ),
        (
            "bullet",
            "Man-Month 계산을 달력 비례·관리자 수동값에서 월별 실제 근무일(평일-공휴일) 자동 산출로 전면 재설계하고 MM=절대공수, 가동률=가용MD 대비 상대지표로 분리해 전 화면 계산식 SSOT 확립",
        ),
        (
            "bullet",
            "프로젝트 Progress 화면을 공수 소진·완성 진척·일정 경과의 3층으로 설계하고 S-curve, 상태 분포, 위험 이슈, 확정/예측 지연과 EVM 해석 가이드를 구현",
        ),
        (
            "bullet",
            "256MB 서버리스 제약에서 전사 Portfolio Dashboard를 scheduled snapshot + read-time aggregation으로 구축하고 매트릭스 조직의 기간 합집합 가동률, 인건비·공수·리스크·T/S/R 지표와 10종 이상의 위젯을 구현",
        ),
        (
            "bullet",
            "분석 화면 5종과 Portfolio 지표를 전수 감사해 미래 기간 포함 가동률 과소, actualMM 20일 고정, 공수=완성 근사, 예산 오표기와 임계값 drift를 교정하고 실제 tunnel E2E에서 scheduled trigger의 asUser 실패도 해결",
        ),
        (
            "bullet",
            "SVAR Gantt 기반 프로젝트 상세 탭에서 일정 drag, 임의 깊이 WBS 계층, Jira native Blocks 선후행 link, 순환 방지·낙관적 update를 구현하고 CSP·readOnly app field·라이선스 제약을 우회",
        ),
        (
            "bullet",
            "가상 folder node·folder/issue mapping·drag-and-drop·삭제 시 reparent·MD 가중 roll-up을 추가하고 Forge SQL의 FK/CASCADE·transaction 부재에 맞춘 자식→부모 멱등 삭제와 순환 이중 방어 구현",
        ),
        (
            "bullet",
            "팀 계층·프로젝트 배정·팀별 분석 탭(차트 6종), 프로젝트→팀→담당자 3-layer 관계도와 팀 일괄 배정을 구현하고 scope 기반 집계·null/0 분모·KST 날짜 경계·손상 parent 처리 오류를 검수",
        ),
        (
            "bullet",
            "팀 관리형 Jira 프로젝트에서 app custom field가 fields 응답에서 누락되는 문제를 실측하고 issue property + jira:entityProperty PoC를 수행; movesRms 단일 property로 coworker·등록일·등록자·WBS를 이전해 회사/팀 관리형 동등성 확보",
        ),
        (
            "bullet",
            "Jira space 구독 시 새 이슈를 RMS 프로젝트에 자동 편입하는 RFC-171을 설계·구현해 이슈 생성 event·시간별 reconcile·선택적 backfill을 연결하고, SCAN→APPLY·IS EMPTY·baseline/watermark 규칙으로 누락과 소유권 충돌 방지",
        ),
        (
            "bullet",
            "PermissionMap·guardedDefine로 109개 resolver를 전수 매핑하고 미매핑 deny-by-default·fail-close·14개 resource key·ALL/MEMBER_ONLY projectScope·PROJECT_COST 분리·SELF accountId 강제를 구현",
        ),
        (
            "bullet",
            "Async Events + Queue 분할 처리와 JobPollingManager로 1회 수천 건 규모 Excel 대량 등록을 안정화하고 portfolio snapshot·issue backfill·설치/업그레이드 migration을 scheduled/event trigger와 연계",
        ),
        (
            "bullet",
            "Forge SQL 1GiB 한도 예방을 위해 information_schema 기반 총량·table별·project별 사용량, 80/90% 경고와 환경별 limit 주입이 가능한 저장공간 관리 화면을 구현",
        ),
        (
            "bullet",
            "팀 관리형 time tracking remainingEstimate 누락 400을 재현해 payload·guard·안내 경로를 수정하고 app field/Jira search index lag·entityProperty date-only·32,768B 제한 등 Jira Cloud 동작을 실측 기반으로 정리",
        ),
        (
            "bullet",
            "page/tab lazy loading과 bundle budget을 도입해 main gzip을 1.73MB에서 271KB로 축소하고 runtime query key 60건을 factory로 수렴; 최신 기능 검증에서 백엔드 1,066개 테스트와 production build 통과",
        ),
        (
            "bullet",
            "제품 소개 16장 source pack, 경쟁 우위 분석, 권한/사용자/기술운영 문서와 배포·도입 산출물 map을 작성해 개발 결과를 영업·운영·사용자 관점의 재사용 가능한 문서로 전환",
        ),

        ("subhead", "[주요 프로젝트/실적]  Elastic Stack 기반 로그 수집·모니터링/검색 플랫폼"),
        ("bullet", "Elasticsearch master/data/ingest 노드 역할과 index·shard·mapping 전략을 설계"),
        (
            "bullet",
            "Metricbeat·Filebeat·Elastic Agent로 서버 및 애플리케이션 로그, Apache/PHP-FPM 성능 지표 수집 pipeline 구축",
        ),
        (
            "bullet",
            "분산 로그를 단일 검색·시각화 환경으로 통합하고 Kibana 운영 지표·통계 dashboard 구성, cluster·pipeline·source 구간별 장애 진단",
        ),

        ("subhead", "[개발 품질·운영 방식]"),
        (
            "bullet",
            "기능마다 설계→BE/FE 구현→계약 대조→자동 테스트→runtime E2E의 closed loop를 운영하고 API 응답·DB side effect·권한 우회 시나리오까지 검증",
        ),
        (
            "bullet",
            "대규모 구조 변경은 worktree·단계별 commit·runtime before/after diff로 격리하고 Flyway migration, Docker image, manifest, CI와 운영 데이터까지 영향 범위를 확인",
        ),
        (
            "bullet",
            "구현 완료·배포 대기·설계 전용 항목을 구분해 문서화하고 미검증 사항은 release gate와 재개 checklist로 남겨 과장 없이 추적",
        ),

        ("subhead", "[이직희망 사유]"),
        (
            "prose",
            "Java/Spring 기반 모니터링 제품 전환과 Zabbix·Elasticsearch 실무 구축, Jira Cloud SaaS의 데이터·권한·운영 문제를 함께 해결해 온 경험을 Elastic/OpenSearch 전문 기술서비스와 엔터프라이즈 검색·AI 플랫폼 영역으로 확장하고, 고객의 문제를 설계부터 배포·운영까지 책임지는 역할을 수행하고자 지원합니다.",
        ),
    ]

    for i in range(start_idx, page_break_idx):
        p = paragraphs[i]._p
        p.getparent().remove(p)

    templates = {
        "subhead": subhead_template,
        "bullet": bullet_template,
        "prose": prose_template,
    }
    parent = page_break_anchor.getparent()
    insert_at = parent.index(page_break_anchor)
    for kind, text in content:
        parent.insert(insert_at, make_paragraph_from_template(templates[kind], text))
        insert_at += 1
    parent.remove(page_break_anchor)


def main() -> None:
    doc = Document(INPUT)
    update_summary_tables(doc)
    insert_core_competency(doc)
    replace_movesai_section(doc)
    doc.core_properties.title = "김찬호 이력서 - 무브스아이 경력 상세 보강본"
    doc.core_properties.subject = "MovesEye 및 Moves Workforce 수행 업무 상세 경력기술"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
